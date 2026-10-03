"""Export Excel : valeurs brutes (nombres, dates), pas de formule injectée, droits de la liste respectés."""

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounting import services as compta
from apps.accounting.models import Journal, SensEcriture
from apps.accounting.services import LigneSaisie
from apps.accounts.models import Role
from apps.accounts.tests.factories import UserFactory
from apps.billing.tests.helpers import JOUR, direction, emise, finances
from apps.core.xlsx import TYPE_XLSX, classeur

pytestmark = pytest.mark.django_db


def _ouvrir(reponse):
    assert reponse.status_code == 200
    assert reponse["Content-Type"] == TYPE_XLSX
    assert reponse["Content-Disposition"].startswith('attachment; filename="')
    return load_workbook(BytesIO(reponse.content))


def _lignes(feuille):
    return [[c.value for c in ligne] for ligne in feuille.iter_rows()]


# --- le classeur ---


def test_les_montants_sont_des_nombres_les_dates_des_dates_et_les_formules_du_texte():
    octets = classeur([{
        "titre": "Essai", "sous_titre": "Période", "entetes": ["A", "B", "C", "D"],
        "lignes": [[Decimal("1234.50"), date(2026, 3, 1), datetime(2026, 3, 1, 8, 30, tzinfo=timezone.utc), "=1+1"]],
        "pied": ["", "", "Total", Decimal("1234.50")],
    }])

    feuille = load_workbook(BytesIO(octets))["Essai"]

    assert feuille["A5"].value == 1234.5 and feuille["A5"].number_format == "#,##0.00"
    assert feuille["B5"].value.date() == date(2026, 3, 1)
    assert feuille["D5"].value == "=1+1" and feuille["D5"].data_type == "s"  # du texte, jamais une formule
    assert feuille["D6"].value == 1234.5 and feuille["D6"].font.bold


def test_les_titres_de_feuille_sont_rendus_valides_et_uniques():
    octets = classeur([
        {"titre": "Bilan: actif/passif [2026]", "entetes": ["X"], "lignes": []},
        {"titre": "Bilan: actif/passif [2026]", "entetes": ["X"], "lignes": []},
    ])

    noms = load_workbook(BytesIO(octets)).sheetnames

    assert len(noms) == 2 and len(set(n.lower() for n in noms)) == 2
    assert all(len(n) <= 31 and not set("[]:*?/\\") & set(n) for n in noms)


# --- listes ---


def test_export_des_factures_avec_montants_numeriques(client):
    emise(prix="1000000")
    client.force_login(finances())

    feuille = _ouvrir(client.get(reverse("billing:factures_xlsx"))).active

    lignes = _lignes(feuille)
    entetes = next(l for l in lignes if l[0] == "N°")
    donnees = lignes[lignes.index(entetes) + 1]
    assert entetes[-4:] == ["HT (FCFA)", "TVA (FCFA)", "TTC (FCFA)", "Reste à recouvrer (FCFA)"]
    assert donnees[-4] == 1000000 and isinstance(donnees[-2], (int, float))


def test_export_des_depenses_et_droits(client):
    from apps.billing import services as billing
    from apps.billing.models import ModePaiement

    billing.enregistrer_depense(
        finances(), categorie="PEAGES", date_depense=JOUR, libelle="=SOMME(A1)", montant=Decimal("11800"),
        mode=ModePaiement.ESPECES, montant_tva=Decimal("1800"),
    )
    client.force_login(finances())

    feuille = _ouvrir(client.get(reverse("billing:depenses_xlsx"))).active
    lignes = _lignes(feuille)

    ligne = next(l for l in lignes if l and l[1] == "=SOMME(A1)")
    assert ligne[4] == 11800 and ligne[5] == 1800
    client.force_login(UserFactory(role=Role.CHAUFFEUR))
    assert client.get(reverse("billing:depenses_xlsx")).status_code == 403


def test_export_de_la_tresorerie_signe_les_sorties(client):
    from apps.billing import services as billing
    from apps.billing.models import ModePaiement

    billing.enregistrer_reglement(
        emise(prix="1000000"), finances(), montant=Decimal("500000"), mode=ModePaiement.VIREMENT, date_reglement=JOUR
    )
    billing.enregistrer_depense(
        finances(), categorie="PEAGES", date_depense=JOUR, libelle="Péage", montant=Decimal("100000"),
        mode=ModePaiement.ESPECES,
    )
    client.force_login(finances())

    feuille = _ouvrir(client.get(reverse("finance:xlsx"), {"date_debut": "2020-01-01", "date_fin": "2099-12-31"})).active
    lignes = [l for l in _lignes(feuille) if l and isinstance(l[0], (datetime, date))]

    montants = sorted(l[5] for l in lignes)
    assert montants == [-100000, 500000]
    assert _lignes(feuille)[-1][-1] == 400000  # total


# --- rapports comptables ---


def _ecritures():
    compta.passer_ecriture(
        journal=Journal.VENTES, date_ecriture=date(2026, 3, 1), libelle="Vente",
        lignes=[
            LigneSaisie(compte="411000", sens=SensEcriture.DEBIT, montant=Decimal("1000")),
            LigneSaisie(compte="706100", sens=SensEcriture.CREDIT, montant=Decimal("1000")),
        ],
    )


def test_balance_en_excel_avec_totaux(client):
    _ecritures()
    client.force_login(finances())

    feuille = _ouvrir(client.get(reverse("accounting:balance_xlsx"))).active
    lignes = _lignes(feuille)

    pied = lignes[-1]
    assert pied[1] == "Total" and pied[2] == 1000 and pied[3] == 1000
    assert any(l[0] == "411000" and l[2] == 1000 for l in lignes)


def test_grand_livre_en_excel_exige_un_compte(client):
    _ecritures()
    client.force_login(finances())
    from apps.accounting.models import Compte

    assert client.get(reverse("accounting:grand_livre_xlsx")).status_code == 404
    compte = Compte.objects.get(numero="411000")

    feuille = _ouvrir(client.get(reverse("accounting:grand_livre_xlsx"), {"compte": compte.numero})).active
    lignes = _lignes(feuille)

    assert any(l[0] is not None and l[4] == 1000 and l[6] == 1000 for l in lignes[4:])


def test_bilan_et_compte_de_resultat_en_excel(client):
    _ecritures()
    client.force_login(direction())

    bilan = _ouvrir(client.get(reverse("accounting:bilan_xlsx")))
    resultat = _ouvrir(client.get(reverse("accounting:compte_resultat_xlsx")))

    assert bilan.sheetnames == ["Actif", "Passif"]
    assert _lignes(bilan["Actif"])[-1] == [None, "Total actif", 1000]
    assert _lignes(bilan["Passif"])[-1] == [None, "Total passif", 1000]  # résultat en cours compris
    assert resultat.sheetnames == ["Produits", "Charges", "Résultat"]
    assert _lignes(resultat["Résultat"])[-1][1] == 1000


def test_declaration_tva_en_excel(client):
    client.force_login(finances())

    feuille = _ouvrir(client.get(reverse("accounting:declaration_tva_xlsx"))).active

    assert [l[0] for l in _lignes(feuille)[-3:]][0].startswith("TVA collectée")


@pytest.mark.parametrize("nom", ["balance_xlsx", "bilan_xlsx", "compte_resultat_xlsx", "declaration_tva_xlsx"])
def test_les_exports_comptables_sont_reserves_aux_roles_de_consultation(client, nom):
    client.force_login(UserFactory(role=Role.PARCAUTO))

    assert client.get(reverse(f"accounting:{nom}")).status_code == 403


@pytest.mark.parametrize(
    "ecran, export, role",
    [
        ("billing:factures", "billing:factures_xlsx", Role.FINANCES),
        ("billing:depenses", "billing:depenses_xlsx", Role.FINANCES),
        ("finance:tresorerie", "finance:xlsx", Role.FINANCES),
        ("accounting:balance", "accounting:balance_xlsx", Role.FINANCES),
        ("accounting:bilan", "accounting:bilan_xlsx", Role.FINANCES),
        ("accounting:compte_resultat", "accounting:compte_resultat_xlsx", Role.FINANCES),
        ("accounting:declaration_tva", "accounting:declaration_tva_xlsx", Role.FINANCES),
        ("accounting:grand_livre", "accounting:grand_livre_xlsx", Role.FINANCES),
    ],
)
def test_chaque_ecran_propose_le_bouton_excel(client, ecran, export, role):
    client.force_login(UserFactory(role=role))

    page = client.get(reverse(ecran)).content.decode()

    assert reverse(export) in page and "fa-file-excel" in page
