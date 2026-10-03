"""Export Excel (.xlsx) des listes et des rapports.

Le même contenu que l'impression, mais exploitable dans un tableur : les montants sont de vrais nombres
(totaux, tris et filtres fonctionnent) et les dates de vraies dates, pas du texte formaté.

Une feuille = un dictionnaire ``{"titre", "sous_titre", "entetes", "lignes", "pied"}`` (``sous_titre`` et ``pied``
facultatifs ; ``pied`` est une ligne de totaux en gras). :func:`reponse_classeur` en fait un téléchargement.

Un texte qui commence par ``=`` est écrit comme du texte, jamais comme une formule : un libellé saisi par un
utilisateur ne doit pas pouvoir s'exécuter à l'ouverture du fichier.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FORMAT_MONTANT = "#,##0.00"
FORMAT_DATE = "dd/mm/yyyy"
FORMAT_DATE_HEURE = "dd/mm/yyyy hh:mm"
LARGEUR_MAX = 60
INTERDITS_TITRE = set("[]:*?/\\")


def _ecrire(cellule, valeur) -> None:
    if isinstance(valeur, bool):
        cellule.value = "Oui" if valeur else "Non"
    elif isinstance(valeur, Decimal):
        cellule.value = float(valeur)
        cellule.number_format = FORMAT_MONTANT
    elif isinstance(valeur, datetime):
        if timezone.is_aware(valeur):
            valeur = timezone.localtime(valeur).replace(tzinfo=None)
        cellule.value = valeur
        cellule.number_format = FORMAT_DATE_HEURE
    elif isinstance(valeur, date):
        cellule.value = valeur
        cellule.number_format = FORMAT_DATE
    elif isinstance(valeur, str):
        cellule.value = valeur
        if valeur.startswith("="):
            cellule.data_type = "s"  # du texte, pas une formule
    else:
        cellule.value = valeur


def _titre_feuille(titre: str, deja: set[str]) -> str:
    propre = "".join("-" if c in INTERDITS_TITRE else c for c in titre).strip()[:31] or "Feuille"
    candidat, n = propre, 2
    while candidat.lower() in deja:
        suffixe = f" ({n})"
        candidat = propre[: 31 - len(suffixe)] + suffixe
        n += 1
    deja.add(candidat.lower())
    return candidat


def classeur(feuilles: list[dict]) -> bytes:
    """Classeur Excel (octets) : une feuille par élément de ``feuilles``."""
    wb = Workbook()
    wb.remove(wb.active)
    noms: set[str] = set()
    gras = Font(bold=True)
    fond_entete = PatternFill("solid", fgColor="E2E8F0")
    for donnees in feuilles:
        feuille = wb.create_sheet(_titre_feuille(donnees["titre"], noms))
        ligne = 1
        feuille.cell(row=ligne, column=1, value=donnees["titre"]).font = Font(bold=True, size=14)
        ligne += 1
        if donnees.get("sous_titre"):
            _ecrire(feuille.cell(row=ligne, column=1), donnees["sous_titre"])
            ligne += 1
        ligne += 1
        ligne_entetes = ligne
        for colonne, entete in enumerate(donnees["entetes"], start=1):
            cellule = feuille.cell(row=ligne, column=colonne, value=entete)
            cellule.font, cellule.fill = gras, fond_entete
            cellule.alignment = Alignment(wrap_text=True, vertical="center")
        for valeurs in donnees["lignes"]:
            ligne += 1
            for colonne, valeur in enumerate(valeurs, start=1):
                _ecrire(feuille.cell(row=ligne, column=colonne), valeur)
        if donnees.get("pied"):
            ligne += 1
            for colonne, valeur in enumerate(donnees["pied"], start=1):
                cellule = feuille.cell(row=ligne, column=colonne)
                _ecrire(cellule, valeur)
                cellule.font = gras
        feuille.freeze_panes = feuille.cell(row=ligne_entetes + 1, column=1)
        for colonne, entete in enumerate(donnees["entetes"], start=1):
            plus_long = max(
                [len(str(entete))]
                + [len(str(v)) for v in (ligne_[colonne - 1] for ligne_ in donnees["lignes"] if len(ligne_) >= colonne)]
            )
            feuille.column_dimensions[get_column_letter(colonne)].width = min(max(plus_long + 2, 10), LARGEUR_MAX)
    tampon = BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def reponse_classeur(nom_fichier: str, feuilles: list[dict]) -> HttpResponse:
    """Téléchargement d'un classeur ; ``nom_fichier`` sans extension (la date du jour y est ajoutée)."""
    reponse = HttpResponse(classeur(feuilles), content_type=TYPE_XLSX)
    reponse["Content-Disposition"] = f'attachment; filename="{nom_fichier}_{timezone.localdate():%Y-%m-%d}.xlsx"'
    return reponse


class ExportXlsxMixin:
    """Fait d'une vue d'impression de liste (:class:`apps.core.views.ImpressionListeMixin`) un export Excel.

    ``class XExporterXlsxView(ExportXlsxMixin, XImprimerView)`` : mêmes droits, mêmes filtres, mêmes lignes
    que l'impression. ``colonnes`` y est redéclaré avec des **valeurs brutes** (nombres, dates) au lieu du texte
    formaté de l'impression ; ``nom_fichier`` donne le nom du téléchargement.
    """

    nom_fichier = "export"
    limite_impression = 10000  # un tableur supporte bien plus qu'une page imprimée

    def render_to_response(self, context, **kwargs):
        feuille = {
            "titre": self.get_titre_impression(),
            "sous_titre": self.get_sous_titre_impression(),
            "entetes": context["entetes"],
            "lignes": context["lignes"],
        }
        if context.get("tronque"):
            feuille["pied"] = [f"Export limité aux {self.limite_impression} premières lignes : affinez les filtres."]
        return reponse_classeur(self.nom_fichier, [feuille])
