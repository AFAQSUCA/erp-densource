"""Import des données de l'entreprise depuis Excel : une feuille par type de donnée, tout ou rien, rien d'écrasé."""

from io import BytesIO
from pathlib import Path

import pytest
from django.conf import settings
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import Role
from apps.accounts.tests.factories import UserFactory
from apps.audit.models import ActionChoices, AuditLog
from apps.customers.models import Client
from apps.drivers.models import Chauffeur, Copilote, StatutChauffeur
from apps.fleet.models import DocumentReglementaire, StatutVehicule, TypeDocument, Vehicule
from apps.hr.models import Personnel
from apps.importation import modele, services

from .helpers import (
    CHAUFFEUR, CHAUFFEUR_PERSONNE, CLIENT, COPILOTE, COPILOTE_PERSONNE, PERSONNEL, VEHICULE, classeur, modele_vide_lu,
)

pytestmark = pytest.mark.django_db


def _admin():
    return UserFactory(role=Role.ADMIN)


def _importer(fichier, **kwargs):
    return services.importer_classeur(fichier, _admin(), **kwargs)


def _tout():
    return classeur(
        personnel=[PERSONNEL, CHAUFFEUR_PERSONNE, COPILOTE_PERSONNE], chauffeurs=[CHAUFFEUR], copilotes=[COPILOTE],
        clients=[CLIENT], vehicules=[[*VEHICULE[:9], "Soro Ali", *VEHICULE[10:]]],
    )


# --- le modèle ---


def test_le_modele_a_les_six_feuilles_et_les_en_tetes_que_l_import_lit():
    wb = modele_vide_lu()

    assert wb.sheetnames == ["Instructions"] + [modele.TITRES[c] for c in (
        modele.PERSONNEL, modele.VEHICULES, modele.CHAUFFEURS, modele.COPILOTES, modele.CLIENTS)]
    for cle in modele.ORDRE:
        entetes = [c.value for c in wb[modele.TITRES[cle]][1]][: len(modele.COLONNES[cle])]
        assert entetes == modele.COLONNES[cle]


def test_les_en_tetes_du_personnel_sont_ceux_de_l_import_rh_existant():
    from apps.hr.services import COLONNES_IMPORT

    assert modele.COLONNES[modele.PERSONNEL] == list(COLONNES_IMPORT)


def test_la_copie_versionnee_du_modele_a_les_memes_feuilles_et_en_tetes():
    chemin = Path(settings.BASE_DIR) / "modele-donnees-entreprise-DEN-Source.xlsx"
    wb = load_workbook(chemin)

    assert wb.sheetnames == modele_vide_lu().sheetnames
    for cle in modele.ORDRE:
        assert [c.value for c in wb[modele.TITRES[cle]][1]][: len(modele.COLONNES[cle])] == modele.COLONNES[cle]


def test_un_modele_non_modifie_s_importe_sans_rien_creer():
    """Le modèle brut (avec ses exemples) est valide, et ses exemples ne deviennent jamais de vraies fiches."""
    fichier = BytesIO(modele.construire_modele())

    rapport = _importer(fichier)

    assert rapport.enregistre and rapport.total_crees == 0
    assert sum(f.exemples_ignores for f in rapport.feuilles) == 5
    assert not Personnel.objects.exists() and not Client.objects.exists() and not Vehicule.objects.exists()


# --- import complet ---


def test_un_import_complet_cree_et_relie_tout():
    rapport = _importer(_tout())

    assert rapport.enregistre and not rapport.erreurs
    assert rapport.total_crees == 3 + 1 + 1  # 3 personnes, 1 client, 1 camion
    assert rapport.total_mis_a_jour == 2  # fiches chauffeur et copilote complétées
    chauffeur = Chauffeur.objects.get(personnel__nom="Soro")
    assert chauffeur.numero_permis == "CI-PL-099999" and chauffeur.categories_permis == ["C", "E"]
    assert str(chauffeur.date_expiration_permis) == "2028-11-12"
    copilote = Copilote.objects.get(personnel__nom="Diomandé")
    assert copilote.telephone == "+225 01 00 00 00 00" and copilote.statut == StatutChauffeur.SUSPENDU
    client = Client.objects.get(ncc_nif="CI-NCC-9876543B")
    assert client.taux_tva == 18 and client.delai_paiement_jours == 30
    camion = Vehicule.objects.get(immatriculation="CI-9999-ZZ")
    assert camion.chauffeur_habituel == chauffeur and camion.kilometrage == 185000
    assert Personnel.objects.get(nom="Diallo").matricule.startswith("PERS-")


def test_les_documents_du_camion_sont_enregistres_avec_une_delivrance_par_defaut():
    _importer(_tout())

    assurance = DocumentReglementaire.objects.get(type_document=TypeDocument.ASSURANCE)
    carte = DocumentReglementaire.objects.get(type_document=TypeDocument.CARTE_GRISE)
    assert (str(assurance.date_delivrance), str(assurance.date_expiration)) == ("2026-01-01", "2027-01-01")
    assert (str(carte.date_delivrance), str(carte.date_expiration)) == ("2026-03-15", "2027-03-15")  # un an avant
    assert DocumentReglementaire.objects.count() == 2


def test_chaque_creation_est_tracee_au_journal_d_audit_avec_l_importateur():
    _importer(_tout())

    assert AuditLog.objects.filter(entite="Personnel", action=ActionChoices.CREATE).count() == 3
    assert AuditLog.objects.filter(entite="Client", action=ActionChoices.CREATE).count() == 1
    assert AuditLog.objects.filter(entite="Vehicule", action=ActionChoices.CREATE).count() == 1


def test_reimporter_le_meme_fichier_ne_cree_aucun_doublon():
    _importer(_tout())

    rapport = _importer(_tout())

    assert rapport.enregistre and rapport.total_crees == 0
    assert sum(f.deja_presents for f in rapport.feuilles) == 3 + 1 + 1
    assert Personnel.objects.count() == 3 and Client.objects.count() == 1 and Vehicule.objects.count() == 1


def test_une_fiche_deja_dans_l_erp_n_est_jamais_ecrasee():
    from apps.customers.tests.factories import ClientFactory

    ClientFactory(ncc_nif="CI-NCC-9876543B", raison_sociale="Nom d'origine")

    rapport = _importer(classeur(clients=[CLIENT]))

    assert rapport.feuilles[3].deja_presents == 1
    assert Client.objects.get(ncc_nif="CI-NCC-9876543B").raison_sociale == "Nom d'origine"


def test_une_feuille_peut_etre_importee_seule():
    rapport = _importer(classeur(clients=[CLIENT]))

    assert rapport.enregistre and rapport.total_crees == 1
    assert [f.presente for f in rapport.feuilles] == [False, False, False, True, False]


def test_les_chauffeurs_se_rattachent_a_une_fiche_deja_dans_l_erp():
    from apps.hr.tests.factories import PersonnelFactory

    PersonnelFactory(nom="Soro", prenom="Ali", poste="Chauffeur")

    rapport = _importer(classeur(chauffeurs=[CHAUFFEUR]))

    assert rapport.enregistre and Chauffeur.objects.get(personnel__nom="Soro").numero_permis == "CI-PL-099999"


def test_les_noms_se_rapprochent_sans_tenir_compte_de_la_casse_ni_des_accents():
    ligne = ["SORO", "ali"] + CHAUFFEUR[2:]
    habituel = [*VEHICULE[:9], "soro ALI", *VEHICULE[10:]]

    rapport = _importer(classeur(personnel=[CHAUFFEUR_PERSONNE], chauffeurs=[ligne], vehicules=[habituel]))

    assert rapport.enregistre and Vehicule.objects.get().chauffeur_habituel.personnel.nom == "Soro"


# --- tout ou rien et erreurs ---


def test_une_ligne_en_erreur_annule_tout_et_la_liste_avec_sa_feuille_et_son_numero():
    mauvais_client = [*CLIENT[:6], 18, "", "abc"]
    mauvais_poste = ["Yao", "Marc", "Astronaute", *PERSONNEL[3:]]

    rapport = _importer(classeur(personnel=[PERSONNEL, mauvais_poste], clients=[mauvais_client]))

    assert not rapport.enregistre
    assert any("Personnel (Employés), ligne 3 : poste « Astronaute » inconnu" in e for e in rapport.erreurs)
    assert any("Clients, ligne 2 : Le délai de paiement" in e for e in rapport.erreurs)
    assert not Personnel.objects.exists() and not Client.objects.exists()  # la ligne 2 valide n'a pas été gardée


@pytest.mark.parametrize(
    "feuille, ligne, attendu",
    [
        ("chauffeurs", ["Inconnu", "Nom"] + CHAUFFEUR[2:], "ne figure ni sur la feuille Personnel"),
        ("chauffeurs", [*CHAUFFEUR[:8], "En mission"], "Le statut « En mission » inconnu"),
        ("chauffeurs", [*CHAUFFEUR[:5], "Z", *CHAUFFEUR[6:]], "Catégorie(s) de permis inconnue(s)"),
        ("copilotes", [*COPILOTE[:4], "En congé"], "Le statut « En congé » inconnu"),
        ("clients", [*CLIENT[:6], 0, "", 30], "motif d'exonération"),
        ("clients", [*CLIENT[:7], "Inventé", 30], "Le motif d'exonération « Inventé » inconnu"),
        ("vehicules", [*VEHICULE[:8], "En mission", ""] + VEHICULE[10:], "Le statut « En mission » inconnu"),
        ("vehicules", [*VEHICULE[:9], "Personne Inconnu", *VEHICULE[10:]], "chauffeur habituel « Personne Inconnu » introuvable"),
        ("vehicules", [*VEHICULE[:3], "deux mille", *VEHICULE[4:9], "", *VEHICULE[10:]], "L'année « deux mille » invalide"),
        ("vehicules", [*VEHICULE[:10], "10/10/2026", "", *VEHICULE[12:]], "Carte grise : la date d'expiration"),
        ("vehicules", [*VEHICULE[:11], "pas une date", *VEHICULE[12:]], "invalide (attendu JJ/MM/AAAA)"),
    ],
)
def test_les_erreurs_de_ligne_sont_expliquees(feuille, ligne, attendu):
    if feuille == "chauffeurs":
        fichier = classeur(personnel=[CHAUFFEUR_PERSONNE], chauffeurs=[ligne])
    elif feuille == "copilotes":
        fichier = classeur(personnel=[COPILOTE_PERSONNE], copilotes=[ligne])
    else:
        fichier = classeur(**{feuille: [ligne]})

    rapport = _importer(fichier)

    assert not rapport.enregistre
    assert any(attendu in e for e in rapport.erreurs), rapport.erreurs


def test_un_chauffeur_dont_le_poste_n_est_pas_chauffeur_est_refuse():
    rapport = _importer(classeur(personnel=[PERSONNEL], chauffeurs=[["Diallo", "Fatou"] + CHAUFFEUR[2:]]))

    assert any("n'a pas de fiche chauffeur" in e for e in rapport.erreurs)


def test_deux_fiches_au_meme_nom_sont_ambigues():
    from apps.hr.tests.factories import PersonnelFactory

    PersonnelFactory(nom="Soro", prenom="Ali", poste="Chauffeur")
    PersonnelFactory(nom="Soro", prenom="Ali", poste="Chauffeur")

    rapport = _importer(classeur(chauffeurs=[CHAUFFEUR]))

    assert any("plusieurs fiches du personnel" in e for e in rapport.erreurs)


def test_des_en_tetes_modifies_refusent_la_feuille():
    wb = modele_vide_lu()
    wb["Clients"]["A1"] = "Société"
    tampon = BytesIO()
    wb.save(tampon)
    tampon.seek(0)

    rapport = _importer(tampon)

    assert not rapport.enregistre and any("en-têtes inattendus" in e for e in rapport.erreurs)


def test_la_simulation_verifie_sans_rien_enregistrer():
    rapport = _importer(_tout(), simulation=True)

    assert not rapport.enregistre and not rapport.erreurs and rapport.total_crees == 5
    assert not Personnel.objects.exists() and not Client.objects.exists() and not Vehicule.objects.exists()


def test_un_camion_au_statut_hors_service_garde_ce_statut():
    ligne = [*VEHICULE[:8], "Hors service", ""] + VEHICULE[10:]

    _importer(classeur(vehicules=[ligne]))

    assert Vehicule.objects.get().statut == StatutVehicule.HORS_SERVICE


def test_les_lignes_vides_sont_ignorees_et_les_fichiers_invalides_refuses():
    assert _importer(classeur(clients=[[], [None] * 9, CLIENT])).total_crees == 1
    with pytest.raises(services.ImportInvalide, match="pas un classeur Excel"):
        _importer(BytesIO(b"ceci n'est pas un classeur"))
    from openpyxl import Workbook

    vide = BytesIO()
    Workbook().save(vide)
    vide.seek(0)
    with pytest.raises(services.ImportInvalide, match="Aucune feuille reconnue"):
        _importer(vide)


def test_les_onglets_acceptent_un_nom_court():
    wb = load_workbook(classeur(clients=[CLIENT]))
    wb.active.title = "clients"
    tampon = BytesIO()
    wb.save(tampon)
    tampon.seek(0)

    assert _importer(tampon).total_crees == 1


# --- droits et écran ---


@pytest.mark.parametrize("role", [Role.DIRECTION, Role.RH, Role.FINANCES, Role.PARCAUTO, Role.CHAUFFEUR])
def test_seul_l_admin_importe(client, role):
    client.force_login(UserFactory(role=role))

    assert client.get(reverse("importation:importer")).status_code == 403
    assert client.get(reverse("importation:modele")).status_code == 403
    assert client.post(reverse("importation:importer"), {"fichier": _tout()}).status_code == 403
    with pytest.raises(services.ImportInvalide, match="administrateur"):
        services.importer_classeur(_tout(), UserFactory(role=role))


def test_le_menu_propose_l_import_a_l_admin_seulement(client):
    client.force_login(_admin())
    assert reverse("importation:importer") in client.get(reverse("home")).content.decode()
    client.force_login(UserFactory(role=Role.RH))
    assert reverse("importation:importer") not in client.get(reverse("home")).content.decode()


def test_l_admin_telecharge_le_modele(client):
    client.force_login(_admin())

    reponse = client.get(reverse("importation:modele"))

    assert reponse.status_code == 200 and "spreadsheetml" in reponse["Content-Type"]
    assert "Clients" in load_workbook(BytesIO(reponse.content)).sheetnames


def test_l_admin_depose_un_fichier_et_voit_le_rapport(client):
    client.force_login(_admin())
    fichier = _tout()
    fichier.name = "donnees.xlsx"

    reponse = client.post(reverse("importation:importer"), {"fichier": fichier})

    page = reponse.content.decode()
    assert reponse.status_code == 200 and "Import terminé" in page and "5 fiches créées" in page
    assert Client.objects.count() == 1


def test_l_ecran_affiche_les_erreurs_par_ligne_et_refuse_un_mauvais_fichier(client):
    client.force_login(_admin())
    mauvais = classeur(clients=[[*CLIENT[:6], 18, "", "abc"]])
    mauvais.name = "donnees.xlsx"

    page = client.post(reverse("importation:importer"), {"fichier": mauvais}).content.decode()
    assert "Rien n'a été importé" in page and "Clients, ligne 2" in page

    faux = BytesIO(b"pas un classeur")
    faux.name = "donnees.xlsx"
    assert "pas un classeur Excel" in client.post(reverse("importation:importer"), {"fichier": faux}).content.decode()
    pdf = BytesIO(b"%PDF")
    pdf.name = "donnees.pdf"
    assert ".xlsx" in client.post(reverse("importation:importer"), {"fichier": pdf}).content.decode()
