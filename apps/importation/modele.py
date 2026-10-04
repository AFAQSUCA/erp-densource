"""Le classeur Excel de collecte des données de l'entreprise : source unique du modèle.

Le même module décrit le modèle téléchargeable (``construire_modele``) et les colonnes que l'import lit
(``services.importer_classeur``) : les deux ne peuvent donc pas diverger. Les listes déroulantes viennent des
énumérations de l'application (postes, départements, statuts...).

Feuilles : Personnel, Véhicules, Chauffeurs, Copilotes, Clients. La feuille « Instructions » n'est pas lue.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from apps.customers.models import MotifExoneration
from apps.fleet.models import StatutVehicule, TypeDocument
from apps.hr.models import POSTES_COURANTS, Departement
from apps.hr.services import COLONNES_IMPORT as COLONNES_PERSONNEL

# Clés des feuilles (ordre de traitement de l'import : un chauffeur doit exister avant de devenir le « chauffeur
# habituel » d'un camion, et une fiche du personnel avant d'être complétée en chauffeur ou copilote).
PERSONNEL, CHAUFFEURS, COPILOTES, CLIENTS, VEHICULES = "personnel", "chauffeurs", "copilotes", "clients", "vehicules"
ORDRE = (PERSONNEL, CHAUFFEURS, COPILOTES, CLIENTS, VEHICULES)

TITRES = {
    PERSONNEL: "Personnel (Employés)",
    VEHICULES: "Véhicules (Camions)",
    CHAUFFEURS: "Chauffeurs",
    COPILOTES: "Copilotes",
    CLIENTS: "Clients",
}
# Mot par lequel commence le nom de l'onglet : « Personnel », « Personnel (Employés) »... sont acceptés.
MOT_CLE = {PERSONNEL: "personnel", VEHICULES: "vehicules", CHAUFFEURS: "chauffeurs", COPILOTES: "copilotes", CLIENTS: "clients"}

# Statuts qui se posent à la main : « En mission » et « En congé » (chauffeurs, copilotes) et « En mission » /
# « En maintenance » (camions) sont posés par les missions, les congés et le garage — jamais depuis un fichier.
STATUTS_PERSONNE = ["Disponible", "Suspendu", "Inactif"]
STATUTS_VEHICULE = [StatutVehicule.DISPONIBLE.label, StatutVehicule.IMMOBILISE.label, StatutVehicule.HORS_SERVICE.label]
CATEGORIES_PERMIS = ["C", "E", "C,E"]
DEPARTEMENTS = [str(libelle) for _, libelle in Departement.choices]
MOTIFS_EXONERATION = [""] + [str(libelle) for _, libelle in MotifExoneration.choices]

DOCUMENTS = [
    (TypeDocument.CARTE_GRISE, "Carte grise"),
    (TypeDocument.ASSURANCE, "Assurance"),
    (TypeDocument.VISITE_TECHNIQUE, "Visite technique"),
    (TypeDocument.PATENTE, "Patente"),
]

COLONNES = {
    PERSONNEL: list(COLONNES_PERSONNEL),
    VEHICULES: [
        "Immatriculation", "Marque", "Modèle", "Année", "N° de châssis (VIN)", "Kilométrage actuel",
        "Capacité de charge (t)", "Réservoir (L)", "Statut", "Chauffeur habituel (Nom Prénom, facultatif)",
    ]
    + [
        libelle
        for _, nom in DOCUMENTS
        for libelle in (f"{nom} — délivrée le (facultatif)", f"{nom} — expire le")
    ],
    CHAUFFEURS: [
        "Nom (comme sur Personnel)", "Prénom (comme sur Personnel)", "Téléphone", "Contact d'urgence", "N° de permis",
        "Catégories de permis", "Expiration du permis", "Expiration visite médicale", "Statut",
    ],
    COPILOTES: ["Nom (comme sur Personnel)", "Prénom (comme sur Personnel)", "Téléphone", "Contact d'urgence", "Statut"],
    CLIENTS: [
        "Raison sociale", "NCC / NIF", "Contact principal", "Téléphone", "Email", "Adresse / siège", "Taux de TVA (%)",
        "Motif d'exonération (si TVA = 0)", "Délai de paiement (jours)",
    ],
}

# Ligne d'exemple (fond jaune) de chaque feuille. L'import l'ignore si elle est restée telle quelle : un exemple
# oublié ne doit jamais devenir une vraie fiche.
EXEMPLES = {
    PERSONNEL: ["Traoré", "Awa", "Comptable", "Ressources Humaines et Finances", "CDI", "01/09/2026", 250000],
    VEHICULES: [
        "CI-1234-AB", "Mercedes-Benz", "Actros", 2019, "WDB9634031L123456", 185000, 25.0, 400, "Disponible",
        "Kouassi Jean", "15/03/2025", "15/03/2027", "01/01/2026", "01/01/2027", "20/06/2026", "20/06/2027",
        "01/01/2026", "31/12/2026",
    ],
    CHAUFFEURS: [
        "Kouassi", "Jean", "+225 07 00 00 00 00", "Kouassi Marie (épouse) +225 05 00 00 00 00", "CI-PL-045678", "C,E",
        "12/11/2028", "03/05/2027", "Disponible",
    ],
    COPILOTES: ["Bamba", "Issouf", "+225 01 00 00 00 00", "Bamba Aïcha (sœur) +225 05 11 11 11 11", "Disponible"],
    CLIENTS: [
        "Cimaf Côte d'Ivoire", "CI-NCC-0123456A", "M. Koné Salif", "+225 27 00 00 00 00", "contact@cimaf.ci",
        "Zone Industrielle, Abidjan", 18.00, "", 30,
    ],
}

INSTRUCTIONS = [
    "DEN Source Group — Classeur de collecte des données de l'entreprise",
    "",
    "Objectif : renseigner vos vraies données (personnel, véhicules, chauffeurs, copilotes, clients) pour les charger "
    "dans l'ERP (menu « Import de données », réservé à l'administrateur). Chaque type de donnée a sa feuille.",
    "",
    "Règles à respecter :",
    "1. Ne modifiez pas la ligne d'en-tête (ligne 1, en gris foncé) de chaque feuille : le libellé et l'ordre des "
    "colonnes doivent rester identiques, sinon l'import refuse la feuille.",
    "2. La ligne d'exemple (fond jaune clair) illustre le format : remplacez-la par vos données ou supprimez-la. Si elle "
    "reste telle quelle, l'import l'ignore.",
    "3. Dates : JJ/MM/AAAA (ex. 05/09/2026), ou une vraie date Excel.",
    "4. Colonnes à liste déroulante : choisissez une valeur dans la liste plutôt que de taper un texte libre.",
    "5. Feuilles « Chauffeurs » et « Copilotes » : indiquez Nom et Prénom EXACTEMENT comme sur la feuille « Personnel » "
    "(ou comme une fiche déjà dans l'ERP), et le poste de la personne doit être « Chauffeur » ou « Copilote ». Le "
    "matricule est généré automatiquement.",
    "6. Statuts : seuls « Disponible », « Suspendu » et « Inactif » (chauffeurs, copilotes) et « Disponible », "
    "« Immobilisé » et « Hors service » (camions) se posent depuis un fichier ; « En mission », « En congé » et « En "
    "maintenance » sont posés par les missions, les congés et le garage.",
    "7. Feuille « Véhicules » : « Chauffeur habituel » est facultatif (Nom Prénom comme sur la fiche du chauffeur). Pour "
    "chaque document, indiquez la date d'expiration ; la date de délivrance est facultative — si elle manque, l'ERP "
    "retient un an avant l'expiration.",
    "8. Laissez une cellule vide quand l'information n'existe pas (n'écrivez pas « N/A »).",
    "9. Rien n'est écrasé : une personne (même Nom et Prénom), un client (même NCC/NIF) ou un camion (même "
    "immatriculation) déjà dans l'ERP est signalé « déjà présent » et laissé tel quel. Les fiches chauffeur et copilote "
    "existantes, elles, sont mises à jour.",
    "10. Tout ou rien : à la moindre ligne en erreur, rien n'est enregistré et chaque erreur est listée avec sa feuille "
    "et son numéro de ligne. La case « Simulation » permet de vérifier le fichier sans rien enregistrer.",
    "11. Le plan comptable (SYSCOHADA) est déjà configuré dans l'ERP : inutile de le renseigner ici.",
]

ENTETE_FILL = PatternFill("solid", fgColor="1F2937")
ENTETE_FONT = Font(color="FFFFFF", bold=True)
EXEMPLE_FILL = PatternFill("solid", fgColor="FEF3C7")
EXEMPLE_FONT = Font(italic=True, color="92400E")
BORDURE = Border(*(Side(style="thin", color="D1D5DB"),) * 4)


def _liste(valeurs: list[str]) -> DataValidation:
    formule = '"' + ",".join(str(v) for v in valeurs) + '"'
    assert len(formule) <= 255, "liste déroulante trop longue pour Excel"
    validation = DataValidation(type="list", formula1=formule, allow_blank=True, showDropDown=False)
    validation.error, validation.errorTitle = "Choisissez une valeur dans la liste.", "Valeur invalide"
    return validation


def _feuille(wb: Workbook, cle: str, largeurs: list[int], listes: dict[str, list[str]]) -> None:
    ws = wb.create_sheet(TITRES[cle])
    for i, texte in enumerate(COLONNES[cle], start=1):
        c = ws.cell(row=1, column=i, value=texte)
        c.fill, c.font, c.border = ENTETE_FILL, ENTETE_FONT, BORDURE
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = largeurs[i - 1]
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    for i, valeur in enumerate(EXEMPLES[cle], start=1):
        c = ws.cell(row=2, column=i, value=valeur)
        c.fill, c.font, c.border = EXEMPLE_FILL, EXEMPLE_FONT, BORDURE
    ws.cell(row=2, column=len(EXEMPLES[cle]) + 2, value="← exemple, à remplacer ou supprimer").font = Font(
        italic=True, color="9CA3AF", size=9
    )
    for lettre, valeurs in listes.items():
        validation = _liste(valeurs)
        ws.add_data_validation(validation)
        validation.add(f"{lettre}2:{lettre}500")


def construire_modele() -> bytes:
    """Le classeur modèle (octets .xlsx) : instructions, puis une feuille par type de donnée."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    ws["A1"] = INSTRUCTIONS[0]
    ws["A1"].font = Font(bold=True, size=14, color="1F2937")
    for i, texte in enumerate(INSTRUCTIONS[1:], start=2):
        c = ws.cell(row=i, column=1, value=texte)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if texte.startswith("Règles"):
            c.font = Font(bold=True, size=11, color="1F2937")
        ws.row_dimensions[i].height = 30
    ws.column_dimensions["A"].width = 110

    _feuille(wb, PERSONNEL, [18, 16, 22, 30, 16, 16, 16], {"C": list(POSTES_COURANTS), "D": DEPARTEMENTS})
    _feuille(
        wb, VEHICULES, [16, 16, 14, 8, 20, 16, 16, 12, 14, 30] + [18] * 8, {"I": STATUTS_VEHICULE}
    )
    _feuille(wb, CHAUFFEURS, [22, 22, 20, 32, 18, 16, 18, 20, 14], {"F": CATEGORIES_PERMIS, "I": STATUTS_PERSONNE})
    _feuille(wb, COPILOTES, [22, 22, 20, 32, 14], {"E": STATUTS_PERSONNE})
    _feuille(wb, CLIENTS, [26, 18, 20, 20, 24, 30, 14, 24, 16], {"H": MOTIFS_EXONERATION})
    tampon = BytesIO()
    wb.save(tampon)
    return tampon.getvalue()
