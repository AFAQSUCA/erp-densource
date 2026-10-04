"""Import des données de l'entreprise depuis le classeur Excel (``modele.py``) — réservé à l'administrateur.

Une seule opération pour les cinq feuilles : personnel, chauffeurs, copilotes, clients, puis camions (dans cet ordre
— un camion peut avoir un chauffeur habituel créé par le même fichier). Chaque ligne passe par le service de son
application (``hr.recruter``, ``customers.creer_client``, ``fleet.creer_vehicule``, ``drivers.modifier_chauffeur``...) :
mêmes contrôles qu'à la saisie, même journal d'audit. Principes :

* **tout ou rien** : à la moindre ligne en erreur, rien n'est enregistré, et chaque erreur est listée avec sa feuille
  et son numéro de ligne ;
* **rien n'est écrasé** : une personne (même nom et prénom), un client (même NCC/NIF) ou un camion (même
  immatriculation) déjà présent est signalé et laissé tel quel — réimporter le même fichier ne crée aucun doublon ;
  seules les fiches chauffeur et copilote existantes sont complétées ;
* **simulation** : le fichier est vérifié de bout en bout puis tout est annulé ;
* une ligne d'exemple restée telle quelle est ignorée.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

import openpyxl
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction

from apps.core.exceptions import ErreurMetier
from apps.core.search import normaliser
from apps.customers import services as clients_services
from apps.customers.exceptions import ClientError
from apps.customers.models import Client, MotifExoneration
from apps.drivers import services as drivers_services
from apps.drivers.exceptions import ChauffeurError
from apps.drivers.models import Chauffeur, Copilote, StatutChauffeur
from apps.fleet import services as fleet_services
from apps.fleet.exceptions import FlotteError
from apps.fleet.models import StatutVehicule, Vehicule
from apps.hr import services as hr_services
from apps.hr.exceptions import PersonnelError
from apps.hr.models import POSTES_COURANTS, Departement, Personnel

from . import modele, permissions

LIGNES_MAX = 2000  # par feuille : au-delà, le fichier n'est pas un classeur de collecte


class ImportInvalide(ErreurMetier):
    """Le fichier lui-même est inutilisable (illisible, aucune feuille reconnue) ou l'acteur n'a pas le droit."""


class _Ligne(Exception):
    """Une ligne qui ne peut pas être importée : le message est affiché tel quel avec son numéro."""


@dataclass
class ResultatFeuille:
    cle: str
    titre: str
    crees: int = 0
    mis_a_jour: int = 0
    deja_presents: int = 0
    exemples_ignores: int = 0
    erreurs: list[str] = field(default_factory=list)
    presente: bool = True


@dataclass
class Rapport:
    feuilles: list[ResultatFeuille]
    simulation: bool
    enregistre: bool

    @property
    def erreurs(self) -> list[str]:
        return [e for f in self.feuilles for e in f.erreurs]

    @property
    def total_crees(self) -> int:
        return sum(f.crees for f in self.feuilles)

    @property
    def total_mis_a_jour(self) -> int:
        return sum(f.mis_a_jour for f in self.feuilles)


# --- lecture des cellules ---


def _texte(valeur) -> str:
    if valeur is None:
        return ""
    if isinstance(valeur, datetime):
        return valeur.strftime("%d/%m/%Y")
    if isinstance(valeur, date):
        return valeur.strftime("%d/%m/%Y")
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    return str(valeur).strip()


def _date(valeur, libelle: str, *, obligatoire: bool = False) -> date | None:
    if valeur in (None, ""):
        if obligatoire:
            raise _Ligne(f"{libelle} est obligatoire.")
        return None
    resultat = hr_services.date_importee(valeur)
    if resultat is None:
        raise _Ligne(f"{libelle} « {valeur} » invalide (attendu JJ/MM/AAAA).")
    return resultat


def _entier(valeur, libelle: str) -> int:
    try:
        nombre = Decimal(_texte(valeur).replace(" ", "").replace(",", "."))
    except InvalidOperation:
        raise _Ligne(f"{libelle} « {valeur} » invalide (nombre entier attendu).") from None
    if nombre != nombre.to_integral_value() or nombre < 0:
        raise _Ligne(f"{libelle} « {valeur} » invalide (nombre entier positif attendu).")
    return int(nombre)


def _decimal(valeur, libelle: str) -> Decimal:
    try:
        return Decimal(_texte(valeur).replace(" ", "").replace(",", "."))
    except InvalidOperation:
        raise _Ligne(f"{libelle} « {valeur} » invalide (nombre attendu).") from None


def _obligatoire(valeur, libelle: str) -> str:
    texte = _texte(valeur)
    if not texte:
        raise _Ligne(f"{libelle} est obligatoire.")
    return texte


def _choix(valeur, correspondances: dict[str, str], libelle: str, *, defaut: str) -> str:
    """Valeur d'une liste fermée (insensible à la casse et aux accents) ; vide = ``defaut``."""
    texte = _texte(valeur)
    if not texte:
        return defaut
    trouve = correspondances.get(normaliser(texte))
    if trouve is None:
        raise _Ligne(f"{libelle} « {texte} » inconnu. Valeurs acceptées : {', '.join(sorted(set(correspondances.values())))}.")
    return trouve


def _cle_nom(nom: str, prenom: str) -> tuple[str, str]:
    return " ".join(normaliser(nom).split()), " ".join(normaliser(prenom).split())


def _index_personnel() -> dict[tuple[str, str], list[Personnel]]:
    index: dict[tuple[str, str], list[Personnel]] = {}
    for personne in Personnel.objects.all():
        index.setdefault(_cle_nom(personne.nom, personne.prenom), []).append(personne)
    return index


def _personne(index, nom: str, prenom: str) -> Personnel:
    trouvees = index.get(_cle_nom(nom, prenom), [])
    if not trouvees:
        raise _Ligne(f"{nom} {prenom} ne figure ni sur la feuille Personnel ni dans l'ERP (même orthographe attendue).")
    if len(trouvees) > 1:
        raise _Ligne(f"{nom} {prenom} : plusieurs fiches du personnel portent ce nom, impossible de choisir.")
    return trouvees[0]


# --- une feuille = une fonction ---


def _personnel(ligne: list, ctx: dict) -> str:
    nom, prenom, poste_brut, dept_brut, contrat, date_brut, salaire_brut = ligne[:7]
    nom, prenom = _obligatoire(nom, "Le nom"), _obligatoire(prenom, "Le prénom")
    if ctx["index"].get(_cle_nom(nom, prenom)):
        return "present"
    poste = hr_services.poste_importe(poste_brut)
    if poste is None:
        raise _Ligne(f"poste « {poste_brut} » inconnu. Postes acceptés : {', '.join(POSTES_COURANTS)}.")
    departement = hr_services.departement_importe(dept_brut)
    if departement is None:
        raise _Ligne(
            f"département « {dept_brut} » inconnu. Acceptés : {', '.join(str(l) for _, l in Departement.choices)}."
        )
    date_embauche = _date(date_brut, "La date d'embauche", obligatoire=True)
    salaire = _decimal(salaire_brut, "Le salaire de base")
    if salaire < 0:
        raise _Ligne(f"le salaire de base « {salaire_brut} » ne peut pas être négatif.")
    personne = hr_services.recruter(
        nom=nom, prenom=prenom, poste=poste, departement=departement, type_contrat=_texte(contrat),
        date_embauche=date_embauche, salaire_base=salaire,
    )
    ctx["index"].setdefault(_cle_nom(nom, prenom), []).append(personne)
    return "cree"


def _categories(valeur) -> list[str]:
    return [c for c in re.split(r"[,\s;/]+", _texte(valeur).upper()) if c]


def _statut_personne(valeur) -> str:
    correspondances = {normaliser(libelle): code for libelle, code in (
        ("Disponible", StatutChauffeur.DISPONIBLE), ("Suspendu", StatutChauffeur.SUSPENDU), ("Inactif", StatutChauffeur.INACTIF),
    )}
    return _choix(valeur, correspondances, "Le statut", defaut=StatutChauffeur.DISPONIBLE)


def _chauffeur(ligne: list, ctx: dict) -> str:
    nom, prenom, telephone, urgence, permis, categories, exp_permis, exp_visite, statut_brut = ligne[:9]
    personne = _personne(ctx["index"], _obligatoire(nom, "Le nom"), _obligatoire(prenom, "Le prénom"))
    fiche = Chauffeur.objects.filter(personnel=personne).first()
    if fiche is None:
        raise _Ligne(f"{personne} n'a pas de fiche chauffeur : son poste doit être « Chauffeur » (feuille Personnel).")
    statut = _statut_personne(statut_brut)
    drivers_services.modifier_chauffeur(
        fiche, telephone=_texte(telephone), contact_urgence=_texte(urgence), numero_permis=_texte(permis),
        categories_permis=_categories(categories),
        date_expiration_permis=_date(exp_permis, "L'expiration du permis"),
        date_expiration_visite_medicale=_date(exp_visite, "L'expiration de la visite médicale"),
    )
    if statut != fiche.statut:
        drivers_services.changer_statut_manuel(fiche, statut)
    return "maj"


def _copilote(ligne: list, ctx: dict) -> str:
    nom, prenom, telephone, urgence, statut_brut = ligne[:5]
    personne = _personne(ctx["index"], _obligatoire(nom, "Le nom"), _obligatoire(prenom, "Le prénom"))
    fiche = Copilote.objects.filter(personnel=personne).first()
    if fiche is None:
        raise _Ligne(f"{personne} n'a pas de fiche copilote : son poste doit être « Copilote » (feuille Personnel).")
    fiche.telephone, fiche.contact_urgence = _texte(telephone), _texte(urgence)
    fiche.save(update_fields=["telephone", "contact_urgence", "updated_at"])
    statut = _statut_personne(statut_brut)
    if statut != fiche.statut:
        drivers_services.changer_statut_copilote(fiche, statut)
    return "maj"


def _client(ligne: list, ctx: dict) -> str:
    raison, ncc, contact, telephone, email, adresse, taux_brut, motif_brut, delai_brut = ligne[:9]
    raison, ncc = _obligatoire(raison, "La raison sociale"), _obligatoire(ncc, "Le NCC / NIF")
    if Client.all_objects.filter(ncc_nif=ncc).exists():
        return "present"
    motifs = {normaliser(str(libelle)): code for code, libelle in MotifExoneration.choices}
    motifs.update({normaliser(code): code for code, _ in MotifExoneration.choices})
    taux = _decimal(taux_brut, "Le taux de TVA") if _texte(taux_brut) else None
    kwargs = dict(
        raison_sociale=raison, ncc_nif=ncc, contact_principal=_obligatoire(contact, "Le contact principal"),
        telephone=_obligatoire(telephone, "Le téléphone"), email=_texte(email), adresse=_obligatoire(adresse, "L'adresse"),
        motif_exoneration=_choix(motif_brut, motifs, "Le motif d'exonération", defaut=""),
    )
    if taux is not None:
        kwargs["taux_tva"] = taux
    if _texte(delai_brut):
        kwargs["delai_paiement_jours"] = _entier(delai_brut, "Le délai de paiement")
    clients_services.creer_client(**kwargs)
    return "cree"


def _vehicule(ligne: list, ctx: dict) -> str:
    immat, marque, mod, annee, vin, km, capacite, reservoir, statut_brut, habituel = ligne[:10]
    immat = _obligatoire(immat, "L'immatriculation")
    immat_normalisee = re.sub(r"\s+", "", immat).upper()
    if Vehicule.all_objects.filter(immatriculation__iexact=immat).exists() or Vehicule.all_objects.filter(
        immatriculation=immat_normalisee
    ).exists():
        return "present"
    statuts = {
        normaliser(str(StatutVehicule(code).label)): code
        for code in (StatutVehicule.DISPONIBLE, StatutVehicule.IMMOBILISE, StatutVehicule.HORS_SERVICE)
    }
    statut = _choix(statut_brut, statuts, "Le statut", defaut=StatutVehicule.DISPONIBLE)
    chauffeur = None
    if _texte(habituel):
        chauffeur = _chauffeur_habituel(_texte(habituel), ctx["index"])
    documents = []
    for rang, (type_document, nom_doc) in enumerate(modele.DOCUMENTS):
        delivree, expire = ligne[10 + 2 * rang], ligne[11 + 2 * rang]
        if expire in (None, "") and delivree in (None, ""):
            continue
        expiration = _date(expire, f"{nom_doc} : la date d'expiration", obligatoire=True)
        delivrance = _date(delivree, f"{nom_doc} : la date de délivrance")
        documents.append((type_document, delivrance or _un_an_avant(expiration), expiration))
    vehicule = fleet_services.creer_vehicule(
        immatriculation=immat, marque=_obligatoire(marque, "La marque"), modele=_obligatoire(mod, "Le modèle"),
        annee=_entier(annee, "L'année"), vin=_obligatoire(vin, "Le n° de châssis (VIN)"),
        kilometrage=_entier(km, "Le kilométrage") if _texte(km) else 0,
        capacite_charge_t=_decimal(capacite, "La capacité de charge"), reservoir_l=_entier(reservoir, "Le réservoir"),
        chauffeur_habituel=chauffeur,
    )
    if statut != vehicule.statut:
        fleet_services.definir_statut(vehicule, statut)
    for type_document, delivrance, expiration in documents:
        fleet_services.enregistrer_document(
            vehicule, type_document=type_document, date_delivrance=delivrance, date_expiration=expiration
        )
    return "cree"


def _un_an_avant(expiration: date) -> date:
    try:
        return expiration.replace(year=expiration.year - 1)
    except ValueError:  # 29 février
        return expiration - timedelta(days=365)


def _chauffeur_habituel(texte: str, index) -> Chauffeur:
    cible = " ".join(normaliser(texte).split())
    trouves = [
        fiche
        for fiche in Chauffeur.objects.select_related("personnel")
        if " ".join(normaliser(f"{fiche.personnel.nom} {fiche.personnel.prenom}").split()) == cible
    ]
    if not trouves:
        raise _Ligne(f"chauffeur habituel « {texte} » introuvable (Nom Prénom comme sur la fiche du chauffeur).")
    if len(trouves) > 1:
        raise _Ligne(f"chauffeur habituel « {texte} » : plusieurs chauffeurs portent ce nom.")
    return trouves[0]


TRAITEMENTS = {
    modele.PERSONNEL: _personnel,
    modele.CHAUFFEURS: _chauffeur,
    modele.COPILOTES: _copilote,
    modele.CLIENTS: _client,
    modele.VEHICULES: _vehicule,
}

ERREURS_METIER = (
    ErreurMetier, PersonnelError, ClientError, ChauffeurError, FlotteError, ValidationError, IntegrityError, ValueError,
)


def _feuilles_du_classeur(classeur) -> dict[str, object]:
    trouvees = {}
    for feuille in classeur.worksheets:
        nom = normaliser(feuille.title)
        for cle, mot in modele.MOT_CLE.items():
            if nom.startswith(mot) and cle not in trouvees:
                trouvees[cle] = feuille
    return trouvees


def _entetes_conformes(feuille, cle: str) -> bool:
    attendues = [normaliser(c) for c in modele.COLONNES[cle]]
    lues = next(feuille.iter_rows(min_row=1, max_row=1, values_only=True), ())
    lues = [normaliser(_texte(c)) for c in lues[: len(attendues)]]
    return lues == attendues


def _est_exemple(ligne: list, cle: str) -> bool:
    exemple = [_texte(v) for v in modele.EXEMPLES[cle]]
    return [_texte(v) for v in ligne[: len(exemple)]] == exemple


def importer_classeur(fichier, acteur, *, simulation: bool = False) -> Rapport:
    """Importe les feuilles reconnues du classeur. Tout ou rien ; ``simulation`` vérifie sans rien enregistrer."""
    if acteur.role_effectif not in permissions.IMPORT_DONNEES:
        raise ImportInvalide("Seul l'administrateur peut importer des données.")
    try:
        classeur = openpyxl.load_workbook(fichier, data_only=True, read_only=True)
    except Exception as erreur:  # fichier corrompu, mauvais format, archive illisible...
        raise ImportInvalide("Ce fichier n'est pas un classeur Excel (.xlsx) lisible.") from erreur
    feuilles = _feuilles_du_classeur(classeur)
    if not feuilles:
        raise ImportInvalide(
            "Aucune feuille reconnue : le classeur doit contenir au moins une feuille « Personnel », « Véhicules », "
            "« Chauffeurs », « Copilotes » ou « Clients » (téléchargez le modèle)."
        )

    resultats = [ResultatFeuille(cle=cle, titre=modele.TITRES[cle], presente=cle in feuilles) for cle in modele.ORDRE]
    with transaction.atomic():
        ctx = {"index": _index_personnel()}
        for resultat in resultats:
            if not resultat.presente:
                continue
            feuille = feuilles[resultat.cle]
            if not _entetes_conformes(feuille, resultat.cle):
                resultat.erreurs.append(
                    f"{resultat.titre} : en-têtes inattendus. Gardez la ligne 1 du modèle telle quelle "
                    f"(colonnes : {', '.join(modele.COLONNES[resultat.cle])})."
                )
                continue
            largeur = len(modele.COLONNES[resultat.cle])
            for numero, brut in enumerate(feuille.iter_rows(min_row=2, values_only=True), start=2):
                if numero - 1 > LIGNES_MAX:
                    resultat.erreurs.append(f"{resultat.titre} : plus de {LIGNES_MAX} lignes, import refusé.")
                    break
                ligne = (list(brut) + [None] * largeur)[:largeur]
                if all(c in (None, "") for c in ligne):
                    continue
                if _est_exemple(ligne, resultat.cle):
                    resultat.exemples_ignores += 1
                    continue
                try:
                    with transaction.atomic():
                        issue = TRAITEMENTS[resultat.cle](ligne, ctx)
                except _Ligne as erreur:
                    resultat.erreurs.append(f"{resultat.titre}, ligne {numero} : {erreur}")
                except ERREURS_METIER as erreur:
                    resultat.erreurs.append(f"{resultat.titre}, ligne {numero} : {erreur}")
                else:
                    if issue == "cree":
                        resultat.crees += 1
                    elif issue == "maj":
                        resultat.mis_a_jour += 1
                    else:
                        resultat.deja_presents += 1
        erreurs = any(r.erreurs for r in resultats)
        enregistre = not erreurs and not simulation
        if not enregistre:
            transaction.set_rollback(True)
    classeur.close()
    return Rapport(feuilles=resultats, simulation=simulation, enregistre=enregistre)
