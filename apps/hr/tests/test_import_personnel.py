"""Recrutement en masse depuis un classeur Excel (services.importer_personnel)."""

import io
from datetime import date

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from apps.accounts.models import Role
from apps.accounts.tests.factories import UserFactory
from apps.drivers.models import Chauffeur
from apps.hr import services
from apps.hr.exceptions import ImportPersonnelError
from apps.hr.models import Departement, Personnel

pytestmark = pytest.mark.django_db

LIGNE_VALIDE = ["Traoré", "Awa", "Comptable", "Comptabilité", "CDI", "01/09/2026", "250000"]


def _classeur(lignes, en_tete=None):
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.append(en_tete or services.COLONNES_IMPORT)
    for ligne in lignes:
        feuille.append(ligne)
    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    return tampon


def _connecte(client, role):
    compte = UserFactory(role=role)
    client.force_login(compte)
    return compte


# --- services.importer_personnel ---


def test_importer_cree_les_fiches_valides():
    crees = services.importer_personnel(_classeur([LIGNE_VALIDE]))

    assert len(crees) == 1
    employe = Personnel.objects.get()
    assert (employe.nom, employe.prenom, employe.poste) == ("Traoré", "Awa", "Comptable")
    assert employe.departement == Departement.COMPTABILITE
    assert employe.date_embauche == date(2026, 9, 1)
    assert employe.matricule.startswith("PERS-")


def test_importer_accepte_le_departement_par_code_ou_par_libelle():
    ligne_code = ["A", "B", "Comptable", "COMPTABILITE", "", "01/09/2026", "100000"]

    services.importer_personnel(_classeur([ligne_code]))

    assert Personnel.objects.get().departement == Departement.COMPTABILITE


def test_importer_ignore_les_lignes_vides():
    lignes = [LIGNE_VALIDE, [None] * 7, ["", "", "", "", "", "", ""]]

    crees = services.importer_personnel(_classeur(lignes))

    assert len(crees) == 1


def test_importer_un_chauffeur_cree_sa_fiche_chauffeur():
    ligne = ["Ouattara", "Moussa", "Chauffeur", "Exploitation", "CDI", "01/09/2026", "200000"]

    services.importer_personnel(_classeur([ligne]))

    assert Chauffeur.objects.filter(personnel__nom="Ouattara").exists()


def test_importer_tout_ou_rien_si_une_ligne_est_invalide():
    ligne_invalide = ["Kone", "Ali", "Poste inexistant", "Comptabilité", "", "01/09/2026", "100000"]

    with pytest.raises(ImportPersonnelError) as exc:
        services.importer_personnel(_classeur([LIGNE_VALIDE, ligne_invalide]))

    assert not Personnel.objects.exists()  # même la ligne valide n'a pas été créée
    assert any("Poste inexistant" in erreur for erreur in exc.value.erreurs)


def test_importer_signale_chaque_type_d_erreur():
    ligne = ["", "Ali", "Poste inconnu", "Département inconnu", "", "pas une date", "pas un nombre"]

    with pytest.raises(ImportPersonnelError) as exc:
        services.importer_personnel(_classeur([ligne]))

    erreurs = "\n".join(exc.value.erreurs)
    assert "nom est obligatoire" in erreurs
    assert "Poste inconnu" in erreurs
    assert "Département inconnu" in erreurs
    assert "pas une date" in erreurs
    assert "pas un nombre" in erreurs


def test_importer_refuse_un_salaire_negatif():
    ligne = ["Kone", "Ali", "Comptable", "Comptabilité", "", "01/09/2026", "-100"]

    with pytest.raises(ImportPersonnelError) as exc:
        services.importer_personnel(_classeur([ligne]))

    assert any("salaire" in erreur for erreur in exc.value.erreurs)


def test_importer_refuse_des_en_tetes_inattendus():
    with pytest.raises(ImportPersonnelError) as exc:
        services.importer_personnel(_classeur([LIGNE_VALIDE], en_tete=["A", "B"]))

    assert not Personnel.objects.exists()
    assert "En-têtes" in exc.value.erreurs[0]


def test_importer_accepte_une_date_deja_au_format_date():
    ligne = list(LIGNE_VALIDE)
    ligne[5] = date(2026, 9, 1)

    services.importer_personnel(_classeur([ligne]))

    assert Personnel.objects.get().date_embauche == date(2026, 9, 1)


# --- écran ---


def test_la_page_d_import_est_accessible_a_la_rh(client):
    _connecte(client, Role.RH)

    assert client.get(reverse("hr:personnel_importer")).status_code == 200


def test_importer_un_fichier_valide_redirige_avec_un_message(client):
    _connecte(client, Role.RH)
    fichier = SimpleUploadedFile(
        "personnel.xlsx",
        _classeur([LIGNE_VALIDE]).read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    reponse = client.post(reverse("hr:personnel_importer"), {"fichier": fichier}, follow=True)

    assert reponse.redirect_chain[-1][0] == reverse("hr:personnel_liste")
    assert any("1 employé(s) importé" in str(m) for m in reponse.context["messages"])


def test_importer_un_fichier_invalide_affiche_les_erreurs_sans_rien_creer(client):
    _connecte(client, Role.RH)
    ligne_invalide = ["Kone", "Ali", "Poste inexistant", "Comptabilité", "", "01/09/2026", "100000"]
    fichier = SimpleUploadedFile(
        "personnel.xlsx",
        _classeur([ligne_invalide]).read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    reponse = client.post(reverse("hr:personnel_importer"), {"fichier": fichier})

    assert reponse.status_code == 200
    assert "Poste inexistant" in reponse.content.decode()
    assert not Personnel.objects.exists()


def test_un_fichier_qui_n_est_pas_un_xlsx_est_refuse(client):
    _connecte(client, Role.RH)
    fichier = SimpleUploadedFile("personnel.txt", b"pas un classeur", content_type="text/plain")

    reponse = client.post(reverse("hr:personnel_importer"), {"fichier": fichier})

    assert reponse.status_code == 200
    assert ".xlsx" in reponse.content.decode()
    assert not Personnel.objects.exists()


def test_le_modele_se_telecharge(client):
    _connecte(client, Role.RH)

    reponse = client.get(reverse("hr:personnel_import_modele"))

    assert reponse.status_code == 200
    assert reponse["Content-Disposition"] == "attachment; filename=modele-import-personnel.xlsx"
    classeur = openpyxl.load_workbook(io.BytesIO(reponse.content))
    premiere_ligne = next(classeur.active.iter_rows(min_row=1, max_row=1, values_only=True))
    assert list(premiere_ligne) == services.COLONNES_IMPORT
